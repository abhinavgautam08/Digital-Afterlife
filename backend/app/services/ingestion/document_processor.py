"""
Document processing module for Phase 2.
Extracts searchable text from PDF, DOCX, TXT, and Markdown files.
Preserves original documents unaltered and bounds memory consumption.
"""
from typing import Dict, Any, Optional
import io

from backend.app.core.logging import logger

try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

MAX_EXTRACTED_CHARS = 150_000  # Guardrail against unbounded memory consumption


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extracts text from readable pages of a PDF document up to safety bounds."""
    if not HAS_PYPDF:
        return ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes), strict=False)
        extracted_pages = []
        total_len = 0

        for page in reader.pages:
            text = page.extract_text()
            if text and text.strip():
                extracted_pages.append(text.strip())
                total_len += len(text)
                if total_len >= MAX_EXTRACTED_CHARS:
                    extracted_pages.append("\n[... content truncated for performance ...]")
                    break

        return "\n\n".join(extracted_pages)
    except Exception as e:
        logger.warning(f"PDF text extraction notice: {e}")
        return ""


def extract_docx_text(file_bytes: bytes) -> str:
    """Extracts paragraph and table text from a Microsoft Word (.docx) document up to safety bounds."""
    if not HAS_DOCX:
        return ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        total_len = 0

        for p in doc.paragraphs:
            if p.text and p.text.strip():
                paragraphs.append(p.text.strip())
                total_len += len(p.text)
                if total_len >= MAX_EXTRACTED_CHARS:
                    paragraphs.append("\n[... content truncated for performance ...]")
                    break

        # Also extract table cells if space remains
        if total_len < MAX_EXTRACTED_CHARS:
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))
                        total_len += sum(len(t) for t in row_texts)
                        if total_len >= MAX_EXTRACTED_CHARS:
                            break

        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"DOCX text extraction notice: {e}")
        return ""


def process_document(
    file_bytes: bytes,
    file_format: str
) -> Dict[str, Any]:
    """
    Extracts text content and document statistics with bounded character limits.
    """
    extracted_text: Optional[str] = None
    meta: Dict[str, Any] = {"format": file_format}

    if file_format == "pdf":
        raw_text = extract_pdf_text(file_bytes)
        if raw_text:
            extracted_text = raw_text

    elif file_format == "docx":
        raw_text = extract_docx_text(file_bytes)
        if raw_text:
            extracted_text = raw_text

    elif file_format in ("text", "txt", "md"):
        try:
            full_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                full_text = file_bytes.decode("latin-1")
            except Exception:
                full_text = ""

        if full_text:
            if len(full_text) > MAX_EXTRACTED_CHARS:
                extracted_text = full_text[:MAX_EXTRACTED_CHARS] + "\n\n[... content truncated for performance ...]"
            else:
                extracted_text = full_text

    if extracted_text:
        meta["character_count"] = len(extracted_text)
        meta["word_count"] = len(extracted_text.split())

    return {
        "extracted_text": extracted_text,
        "width": None,
        "height": None,
        "duration_seconds": None,
        "thumbnail_path": None,
        "original_created_at": None,
        "metadata": meta
    }
